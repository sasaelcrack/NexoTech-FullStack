from io import BytesIO
from xml.sax.saxutils import escape

from app import models


def generar_excel_ventas(ventas, db):
    """Genera un libro XLSX en memoria con el resumen de ventas."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as error:
        raise RuntimeError("Falta instalar openpyxl para exportar Excel") from error

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas"
    hoja.merge_cells("A1:P1")
    hoja["A1"] = "NexoTech | Reporte de ventas"
    hoja["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    hoja["A1"].fill = PatternFill("solid", fgColor="1E40AF")
    hoja["A1"].alignment = Alignment(horizontal="center")
    hoja.merge_cells("A2:P2")
    hoja["A2"] = "Soluciones tecnológicas | Valores expresados en pesos colombianos (COP)"
    hoja["A2"].font = Font(italic=True, color="666666")
    hoja["A2"].alignment = Alignment(horizontal="center")
    encabezados = ["Venta", "Cliente", "Correo cliente", "Responsable", "Fecha", "Estado", "Tipo", "Ítem", "Cantidad", "Precio unitario (COP)", "Subtotal ítem (COP)", "Subtotal venta (COP)", "Descuento (COP)", "Base gravable (COP)", "IVA (19%)", "Total venta (COP)"]
    hoja.append([])
    hoja.append(encabezados)
    for celda in hoja[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E40AF")

    for venta in ventas:
        for detalle in venta.detalles:
            if detalle.producto_id:
                tipo = "Producto"
                articulo = db.query(models.Producto).filter(models.Producto.id == detalle.producto_id).first()
            else:
                tipo = "Servicio"
                articulo = db.query(models.Servicio).filter(models.Servicio.id == detalle.servicio_id).first()
            cliente = db.query(models.Usuario).filter(models.Usuario.id == venta.cliente_id).first()
            responsable = db.query(models.Usuario).filter(models.Usuario.id == venta.usuario_id).first() if venta.usuario_id else None
            hoja.append([
                venta.id,
                f"{cliente.nombre} {cliente.apellido}" if cliente else f"Cliente #{venta.cliente_id}",
                cliente.correo if cliente else "",
                f"{responsable.nombre} {responsable.apellido}" if responsable else "Venta en línea",
                venta.fecha.strftime("%Y-%m-%d %H:%M") if venta.fecha else "",
                venta.estado,
                tipo,
                articulo.nombre if articulo else "Ítem no disponible",
                detalle.cantidad,
                float(detalle.precio_unitario),
                float(detalle.subtotal),
                float(venta.subtotal),
                float(getattr(venta, "descuento", 0) or 0),
                float(venta.subtotal - (getattr(venta, "descuento", 0) or 0)),
                float(venta.impuestos),
                float(venta.total),
            ])
            for celda in hoja[hoja.max_row]:
                if celda.data_type == "f":
                    celda.data_type = "s"
    for columna in ("J", "K", "L", "M", "N", "O", "P"):
        for celda in hoja[columna][4:]:
            celda.number_format = '#,##0'
    for columna, ancho in {"A": 10, "B": 24, "C": 30, "D": 24, "E": 18, "F": 14, "G": 12, "H": 32, "I": 12, "J": 22, "K": 20, "L": 20, "M": 18, "N": 20, "O": 15, "P": 18}.items():
        hoja.column_dimensions[columna].width = ancho
    hoja.freeze_panes = "A5"
    hoja.auto_filter.ref = f"A4:P{hoja.max_row}"
    for fila in hoja.iter_rows(min_row=4):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return archivo


def generar_pdf_tabla(titulo: str, encabezados: list[str], filas: list[list[str]], metadata: list[tuple[str, str]] | None = None, filas_resumen: int = 0):
    """Genera un PDF sencillo en memoria para facturas y reportes."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as error:
        raise RuntimeError("Falta instalar reportlab para exportar PDF") from error

    archivo = BytesIO()
    documento = SimpleDocTemplate(archivo, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=40, bottomMargin=36)
    estilos = getSampleStyleSheet()
    estilo_encabezado = ParagraphStyle("FacturaHeader", parent=estilos["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=9, textColor=colors.white)
    estilo_celda = ParagraphStyle("FacturaCell", parent=estilos["Normal"], fontName="Helvetica", fontSize=8, leading=10, textColor=colors.black, wordWrap="LTR")
    datos_tabla = [
        [Paragraph(escape(str(celda)), estilo_encabezado) for celda in encabezados],
        *[[Paragraph(escape(str(celda)), estilo_celda) for celda in fila] for fila in filas],
    ]
    ancho_util = letter[0] - 72
    anchos = [ancho_util * 0.42, ancho_util * 0.14, ancho_util * 0.22, ancho_util * 0.22]
    if len(encabezados) != 4:
        anchos = [ancho_util / len(encabezados)] * len(encabezados)
    tabla = Table(datos_tabla, repeatRows=1, colWidths=anchos)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E40AF")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    if filas_resumen:
        inicio_resumen = max(1, len(datos_tabla) - filas_resumen)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, inicio_resumen), (-1, -1), colors.HexColor("#F3F4F6")),
            ("LINEABOVE", (0, inicio_resumen), (-1, inicio_resumen), 0.7, colors.HexColor("#9CA3AF")),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#1E40AF")),
            ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ]))
    contenido = [Paragraph(escape(titulo), estilos["Title"]), Spacer(1, 0.08 * inch)]
    if metadata:
        estilo_meta = ParagraphStyle("FacturaMeta", parent=estilos["Normal"], fontSize=9, leading=12, textColor=colors.HexColor("#374151"))
        contenido.append(Table([[Paragraph(f"<b>{escape(etiqueta)}:</b> {escape(valor)}", estilo_meta)] for etiqueta, valor in metadata], colWidths=[ancho_util], style=TableStyle([("BOTTOMPADDING", (0, 0), (-1, -1), 3)])))
        contenido.append(Spacer(1, 0.12 * inch))
    contenido.append(tabla)
    documento.build(contenido)
    archivo.seek(0)
    return archivo
